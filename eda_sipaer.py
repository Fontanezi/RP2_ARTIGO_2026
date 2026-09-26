"""Analise exploratoria da BASE_SIPAER (dados abertos CENIPA, RBAC 121/135) e insercao na Secao 4.

Uso: python eda_sipaer.py <BASE_SIPAER.xlsx> <TRABALHO_RP2.docx> <saida.docx>
A BASE_SIPAER.xlsx e gerada por construir_base.py. As figuras vao para <pasta da saida>/figuras_secao4/.
"""
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

BASE, DOC_IN, DOC_OUT = map(Path, sys.argv[1:4])
FIG_DIR = DOC_OUT.parent / "figuras_secao4"
FIG_DIR.mkdir(parents=True, exist_ok=True)

# ====================================================================== dados
b = pd.read_excel(BASE, sheet_name="base")
dic = pd.read_excel(BASE, sheet_name="dicionario")
fonte = pd.read_excel(BASE, sheet_name="fonte").set_index("item")["valor"]
etapas = [v.rsplit(": ", 1) for k, v in fonte.items() if str(k).startswith("Etapa")]
n = len(b)
b["data_ocorrencia"] = pd.to_datetime(b["data_ocorrencia"])
ini, fim = b.data_ocorrencia.min(), b.data_ocorrencia.max()
acid = b[b.Y == 1]
ig = b[b.classificacao == "INCIDENTE GRAVE"]
inc = b[b.classificacao == "INCIDENTE"]
demais = b[b.Y == 0]
b["local"] = b.aerodromo.map(lambda a: "FORA DE AERÓDROMO" if a == "FAER"
                             else ("AERÓDROMOS SB" if str(a).startswith("SB") else "DEMAIS AERÓDROMOS"))


def pct(a, t=n, casas=1):
    return f"{100 * a / t:.{casas}f}".replace(".", ",")


def br(x, casas=1):
    return f"{x:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def mil(x):
    return f"{int(x):,}".replace(",", ".")


def taxa(col, ordem=None):
    g = b.groupby(col).agg(n=("Y", "size"), a=("Y", "sum"))
    g["tx"] = 100 * g.a / g.n
    return g.reindex(ordem) if ordem else g.sort_values("n", ascending=False)


op = taxa("tipo_operacao", ["REGULAR", "NÃO REGULAR", "TÁXI AÉREO"])
mot = taxa("tipo_motor", ["JATO", "TURBOÉLICE", "TURBOEIXO", "PISTÃO"])
eqp = taxa("tipo_equipamento", ["AVIÃO", "HELICÓPTERO"])
qtd = taxa("qtd_motores", ["BIMOTOR", "MONOMOTOR"])
fase = taxa("fase_voo_grupo", ["SOLO", "DECOLAGEM E SUBIDA", "CRUZEIRO E DESCIDA", "APROXIMAÇÃO E POUSO"])
fab = taxa("fabricante_grupo")
reg = taxa("regiao", ["NORTE", "CENTRO-OESTE", "NORDESTE", "SUL", "SUDESTE"])
loc = taxa("local", ["DEMAIS AERÓDROMOS", "FORA DE AERÓDROMO", "AERÓDROMOS SB"])
ORDEM_CAT = ["FALHA TÉCNICA", "COLISÃO COM FAUNA", "PISTA, POUSO E SOLO", "TRÁFEGO AÉREO", "METEOROLOGIA",
             "PERDA DE CONTROLE OU COLISÃO EM VOO", "OUTROS"]
cat = b.groupby("categoria_tipo").agg(n=("Y", "size"), ig=("classificacao", lambda s: (s == "INCIDENTE GRAVE").sum()),
                                      a=("Y", "sum")).reindex(ORDEM_CAT)
pre, pos = b[b.ano < 2023], b[b.ano >= 2023]
taxi = b[b.tipo_operacao == "TÁXI AÉREO"]
nulos = dic.set_index("coluna")["nulos_%"]
multi = int((b.n_aeronaves_comerciais > 1).sum())
fatais = acid[acid.fatalidades > 0]
fat_acid = acid[acid.n_fatores > 0]
top_fat = fat_acid.fatores_lista.str.split(" | ", regex=False).explode().value_counts()
tl = b.tipos_ocorrencia.fillna("")
tipos_unicos = tl.str.split(" | ", regex=False).explode().replace("", pd.NA).dropna().nunique()
motor_voo = b[(b.categoria_tipo == "FALHA TÉCNICA") & tl.str.contains("FALHA DO MOTOR EM VOO", regex=False)]
excursao = b[tl.str.contains("EXCURSÃO DE PISTA", regex=False)]

# ==================================================================== figuras
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
AZUL = "#1f3b73"

# Figura 1: ocorrencias por ano e tipo de operacao; acidentes por ano
anual = pd.crosstab(b.ano, b.tipo_operacao).reindex(columns=["REGULAR", "NÃO REGULAR", "TÁXI AÉREO"], fill_value=0)
acid_ano = pd.crosstab(acid.ano, acid.tipo_operacao).reindex(index=anual.index,
                                                              columns=anual.columns, fill_value=0)
cores = {"REGULAR": AZUL, "NÃO REGULAR": "#8faadc", "TÁXI AÉREO": "#e39a4c"}
rotulo = {"REGULAR": "Regular", "NÃO REGULAR": "Não regular", "TÁXI AÉREO": "Táxi aéreo"}
fig, (ax, ax2) = plt.subplots(2, 1, figsize=(7.2, 4.4), sharex=True, gridspec_kw={"height_ratios": [2.2, 1]})
for eixo, dados in [(ax, anual), (ax2, acid_ano)]:
    base = pd.Series(0, index=dados.index)
    for c in dados.columns:
        eixo.bar(dados.index, dados[c], bottom=base, color=cores[c], label=rotulo[c], width=0.7)
        base += dados[c]
    for a, tot in base.items():
        eixo.text(a, tot + base.max() * 0.02, str(int(tot)), ha="center", fontsize=7.5)
    eixo.set_ylim(0, base.max() * 1.18)
    eixo.spines[["top", "right"]].set_visible(False)
ax.set_ylabel("Ocorrências")
ax2.set_ylabel("Acidentes")
ax.legend(fontsize=7.5, frameon=False, loc="upper left")
ax2.set_xticks(anual.index)
ax2.set_xticklabels([f"{a}*" if a in (ini.year, fim.year) else str(a) for a in anual.index])
fig.tight_layout()
fig.savefig(FIG_DIR / "fig1_anual.png", dpi=220)
plt.close(fig)

# Figura 2: taxa de acidentes por regiao e por local
fig, (a1, a2) = plt.subplots(1, 2, figsize=(7.2, 2.9), gridspec_kw={"width_ratios": [1.3, 1]})
nomes = {"NORTE": "Norte", "CENTRO-OESTE": "Centro-\nOeste", "NORDESTE": "Nordeste", "SUL": "Sul", "SUDESTE": "Sudeste",
         "DEMAIS AERÓDROMOS": "Demais\naeródromos", "FORA DE AERÓDROMO": "Fora de\naeródromo",
         "AERÓDROMOS SB": "Aeródromos\nSB"}
ymax = max(reg.tx.max(), loc.tx.max()) * 1.2
for a, g, titulo in [(a1, reg, "(a) Região"), (a2, loc, "(b) Local da ocorrência")]:
    a.bar(range(len(g)), g.tx, color=AZUL)
    for i, (_, r) in enumerate(g.iterrows()):
        a.text(i, r.tx + ymax * 0.02, f"{int(r.a)}/{int(r.n)}", ha="center", fontsize=7.5)
    a.set_xticks(range(len(g)))
    a.set_xticklabels([nomes[x] for x in g.index], fontsize=8)
    a.set_ylabel("Acidentes por 100 ocorrências")
    a.set_ylim(0, ymax)
    a.set_title(titulo, fontsize=9)
    a.spines[["top", "right"]].set_visible(False)
fig.tight_layout()
fig.savefig(FIG_DIR / "fig2_regiao_local.png", dpi=220)
plt.close(fig)

# ================================================================== documento
doc = Document(DOC_IN)
paras = doc.paragraphs
i4 = next(i for i, p in enumerate(paras) if p.text.strip().startswith("4. Resultados Parciais"))
i5 = next(i for i, p in enumerate(paras) if p.text.strip().startswith("5. Cronograma"))
assert all(not p.text.strip() for p in paras[i4 + 1:i5]), "Secao 4 do documento de entrada nao esta vazia"
for p in paras[i4 + 1:i5]:
    p._p.getparent().remove(p._p)
ancora = doc.paragraphs[i4 + 1]
assert ancora.text.strip().startswith("5. Cronograma")


def _fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=0, depois=6):
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = 1.15
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)


def paragrafo(texto, tam=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=0, depois=6):
    """Texto com trechos em negrito delimitados por §§."""
    p = ancora.insert_paragraph_before()
    _fmt(p, align, antes, depois)
    for k, parte in enumerate(texto.split("§§")):
        if parte:
            r = p.add_run(parte)
            r.font.size = Pt(tam)
            r.bold = k % 2 == 1
    return p


def subtitulo(texto):
    paragrafo(texto, antes=6, depois=4)


def legenda(rotulo_, texto, antes=2, depois=10):
    p = paragrafo(f"§§{rotulo_}§§ {texto}", tam=10, align=WD_ALIGN_PARAGRAPH.CENTER, antes=antes, depois=depois)
    p.paragraph_format.keep_with_next = rotulo_.startswith("Tabela")


def figura(caminho):
    p = ancora.insert_paragraph_before()
    _fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 6, 0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(caminho), width=Inches(6.0))


def tabela(cabecalho, linhas, larguras, negrito_primeira_col=()):
    t = doc.add_table(rows=1 + len(linhas), cols=len(cabecalho))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "0"), ("color", "000000")):
            e.set(qn(f"w:{k}"), v)
        bordas.append(e)
    t._tbl.tblPr.append(bordas)
    for j, w in enumerate(larguras):
        for cel in t.columns[j].cells:
            cel.width = Inches(w)
    for i, linha in enumerate([cabecalho] + linhas):
        for j, val in enumerate(linha):
            cp = t.cell(i, j).paragraphs[0]
            _fmt(cp, WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER, 1, 1)
            cp.paragraph_format.line_spacing = 1.0
            r = cp.add_run(str(val))
            r.font.size = Pt(10)
            r.bold = i == 0 or linha[0] in negrito_primeira_col
    ancora._p.addprevious(t._tbl)
    paragrafo("", depois=0)


def linha_taxa(nome, r):
    return [nome, mil(r.n), int(r.a), pct(r.a, r.n, 2)]


# ---------------------------------------------------------------- 2.2
import copy

SECAO_22 = [
    "A etapa inicial de preparação consistiu no cruzamento relacional das tabelas de dados abertos do CENIPA "
    "(Ocorrências, Tipos de Ocorrência, Aeronaves, Fatores Contribuintes e Recomendações) pelo código da "
    "ocorrência, resultando em uma linha por ocorrência com os atributos da aeronave comercial envolvida. As "
    "entradas duplicadas foram eliminadas, e os marcadores de ausência usados pelo CENIPA foram convertidos em "
    "valores nulos. Variáveis com elevada taxa de omissão, como a data de publicação do relatório final, foram "
    "descartadas; nas demais, os valores faltantes foram substituídos pela moda, nos atributos categóricos, ou "
    "pela mediana, nos numéricos, calculadas apenas no conjunto de treino para evitar vazamento de informação "
    "para o teste.",
    "A variável dependente (Y) foi definida de forma binária: Y=1 para Acidente e Y=0 para Incidente ou "
    "Incidente Grave. Os preditores restringem-se às informações disponíveis no momento da ocorrência: tipo de "
    "operação, fase do voo, região geográfica, fabricante, tipo e quantidade de motores, tipo de equipamento, "
    "idade e peso máximo de decolagem da aeronave, categoria do tipo de ocorrência, indicador de saída de pista, "
    "mês e dia da semana. Os atributos categóricos foram convertidos por meio de codificação binária ou "
    "codificação pelo alvo, conforme o número de categorias da variável; a codificação pelo alvo é ajustada "
    "dentro de cada dobra da validação cruzada.",
    "Os fatores contribuintes (operacionais, humanos e materiais), as recomendações de segurança, o nível de "
    "dano, as fatalidades e a situação da investigação foram incluídos na base, mas não como preditores. Esses "
    "atributos são produzidos pela investigação, conduzida quase apenas em acidentes e incidentes graves, e por "
    "isso revelariam a própria classe ao modelo (Seção 4.2); são utilizados somente na análise descritiva. O "
    "ano da ocorrência também não é usado como preditor, em razão da mudança na cobertura de notificação a "
    "partir de 2023 (Seção 4.5). Para mitigar o desbalanceamento inerente entre acidentes e "
    "incidentes, aplicou-se a técnica de sobreamostragem sintética da classe minoritária (SMOTE), restrita aos "
    "dados de treino, combinada à compensação por ponderação de pesos das classes durante o treinamento.",
]


def _definir_texto(p, texto):
    runs = p.runs
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


p_a = next(p for p in doc.paragraphs if p.text.startswith("A etapa inicial de preparação"))
p_b = next(p for p in doc.paragraphs if p.text.startswith("A variável dependente (Y)"))
_definir_texto(p_a, SECAO_22[0])
_definir_texto(p_b, SECAO_22[1])
novo = copy.deepcopy(p_b._p)
p_b._p.addnext(novo)
_definir_texto(next(p for p in doc.paragraphs if p._p is novo), SECAO_22[2])

# ---------------------------------------------------------------- 4.1
subtitulo("4.1. Construção e caracterização da base")
paragrafo(
    f"A base foi construída a partir dos dados abertos de ocorrências aeronáuticas do CENIPA, cruzando pelo "
    f"código da ocorrência as tabelas de Ocorrências, Tipos de Ocorrência, Aeronaves, Fatores Contribuintes e "
    f"Recomendações. A versão utilizada contém registros até {fim:%d/%m/%Y}, e a janela de análise cobre os dez "
    f"anos anteriores ({ini:%d/%m/%Y} a {fim:%d/%m/%Y}). O recorte da Seção 2.1 foi aplicado pelo tipo de "
    f"operação da aeronave: foram mantidas as ocorrências com ao menos uma aeronave em operação regular, não "
    f"regular ou de táxi aéreo, que correspondem ao transporte aéreo comercial dos RBAC 121 e 135, e excluídas a "
    f"aviação geral, a instrução, a agrícola e a experimental. A Tabela 1 resume as etapas."
)
legenda("Tabela 1.", "Etapas de construção da base", antes=6, depois=2)
tabela(["Etapa", "Ocorrências"], [[nome, mil(qtd)] for nome, qtd in etapas], [4.6, 1.4])
paragrafo(
    f"Cada linha corresponde a uma ocorrência, com os atributos da aeronave comercial envolvida; nas {multi} "
    f"ocorrências com duas aeronaves comerciais, manteve-se a primeira listada. As tabelas de origem não "
    f"apresentaram linhas duplicadas, e o código da ocorrência é único na base final. Os marcadores de ausência "
    f"usados pelo CENIPA (\"***\", \"NULL\", \"INDETERMINADO\", \"DESCONHECIDO\") foram convertidos em valores "
    f"nulos. A omissão é inferior a 1% em todos os atributos categóricos candidatos a preditor e mais alta "
    f"apenas no número de assentos ({br(nulos['assentos'])}%) e no ano de fabricação "
    f"({br(nulos['ano_fabricacao'])}%). Como a imputação descrita na Seção 2.2 é ajustada somente no "
    f"conjunto de treino, os nulos foram mantidos na "
    f"base publicada. O arquivo final (BASE_SIPAER.xlsx) tem {mil(n)} linhas e {b.shape[1] - 1} colunas e acompanha um "
    f"dicionário de dados que indica a origem de cada coluna e se ela pode ser usada como preditor."
)

# ---------------------------------------------------------------- 4.2
subtitulo("4.2. Variável alvo e atributos posteriores ao evento")
paragrafo(
    f"Com Y = 1 para acidente e Y = 0 para incidente ou incidente grave, a base contém {len(acid)} acidentes "
    f"({pct(len(acid), casas=2)}%), {len(ig)} incidentes graves ({pct(len(ig))}%) e {mil(len(inc))} incidentes "
    f"({pct(len(inc))}%), cerca de um acidente para cada {round(n / len(acid))} ocorrências. Dos acidentes, "
    f"{len(fatais)} tiveram vítimas fatais, somando {int(acid.fatalidades.sum())} mortes, "
    f"{int(acid.fatalidades.max())} delas no acidente de Vinhedo (SP), em 09/08/2024. Na partição estratificada "
    f"80/20, o conjunto de teste receberá cerca de {round(len(acid) * 0.2)} acidentes, e cada dobra da validação "
    f"cruzada com k = 5, um número semelhante, o que é suficiente para as métricas da Seção 2.4, embora com "
    f"variância considerável."
)
leak = [
    ("Dano substancial ou aeronave destruída", b.nivel_dano.isin(["SUBSTANCIAL", "DESTRUÍDA"])),
    ("Relatório final publicado", b.relatorio_publicado.eq("SIM")),
    ("Com fatores contribuintes registrados", b.n_fatores > 0),
    ("Com recomendações de segurança", b.n_recomendacoes > 0),
    ("Investigação ainda ativa", b.status_investigacao.eq("ATIVA")),
]
paragrafo(
    "Vários atributos das tabelas cruzadas só existem porque a ocorrência foi investigada, e a investigação "
    "completa é feita quase apenas para acidentes e incidentes graves. A Tabela 2 mostra o efeito: os fatores "
    "contribuintes e as recomendações, obtidos no cruzamento descrito na Seção 2.2, aparecem na maioria dos acidentes e "
    "em uma fração ínfima das demais ocorrências. Usá-los como preditores entregaria ao modelo a própria "
    "resposta, e por isso a Seção 2.2 os exclui dos preditores. O mesmo vale para o nível de dano, que integra a definição de acidente, para as fatalidades e "
    "para a situação da investigação. Esses atributos permanecem na base para fins descritivos e estão "
    "marcados no dicionário como posteriores ao evento."
)
legenda("Tabela 2.", "Atributos produzidos pela investigação, por classe", antes=6, depois=2)
tabela(["Atributo", f"Acidentes (n = {len(acid)})", f"Demais (n = {mil(len(demais))})"],
       [[nome, f"{int(m[b.Y == 1].sum())} ({pct(m[b.Y == 1].sum(), len(acid))}%)",
         f"{int(m[b.Y == 0].sum())} ({pct(m[b.Y == 0].sum(), len(demais))}%)"] for nome, m in leak],
       [3.2, 1.45, 1.45])
paragrafo(
    f"Entre os {len(fat_acid)} acidentes com fatores registrados, {int(fat_acid.fator_operacional.sum())} têm "
    f"fatores da área operacional, {int(fat_acid.fator_humano.sum())} da área humana e "
    f"{int(fat_acid.fator_material.sum())} da área material. Os mais frequentes são processo decisório "
    f"({top_fat.get('PROCESSO DECISÓRIO', 0)}), julgamento de pilotagem ({top_fat.get('JULGAMENTO DE PILOTAGEM', 0)}) "
    f"e supervisão gerencial ({top_fat.get('SUPERVISÃO GERENCIAL', 0)}). Esses resultados caracterizam os "
    f"acidentes já investigados, mas não servem para antecipar a severidade de uma ocorrência nova."
)

# ---------------------------------------------------------------- 4.3
subtitulo("4.3. Operação e características da aeronave")
paragrafo(
    f"O tipo de operação é o atributo que mais separa as classes (Tabela 3). A operação regular responde por "
    f"{pct(op.loc['REGULAR', 'n'])}% das ocorrências e por apenas {int(op.loc['REGULAR', 'a'])} acidentes "
    f"({pct(op.loc['REGULAR', 'a'], op.loc['REGULAR', 'n'], 2)}%), enquanto o táxi aéreo, com "
    f"{pct(op.loc['TÁXI AÉREO', 'n'])}% das ocorrências, concentra {int(op.loc['TÁXI AÉREO', 'a'])} dos "
    f"{len(acid)} acidentes ({pct(op.loc['TÁXI AÉREO', 'a'], op.loc['TÁXI AÉREO', 'n'])}%). As características "
    f"da aeronave seguem o mesmo padrão: a taxa de acidentes é de {pct(mot.loc['PISTÃO', 'a'], mot.loc['PISTÃO', 'n'])}% "
    f"em aeronaves a pistão contra {pct(mot.loc['JATO', 'a'], mot.loc['JATO', 'n'], 2)}% nos jatos, de "
    f"{pct(qtd.loc['MONOMOTOR', 'a'], qtd.loc['MONOMOTOR', 'n'])}% nas monomotoras e de "
    f"{pct(eqp.loc['HELICÓPTERO', 'a'], eqp.loc['HELICÓPTERO', 'n'])}% nos helicópteros. Entre os fabricantes, "
    f"Airbus e Boeing somam {mil(fab.loc['AIRBUS', 'n'] + fab.loc['BOEING', 'n'])} ocorrências e "
    f"{int(fab.loc['AIRBUS', 'a'] + fab.loc['BOEING', 'a'])} acidentes, enquanto Cessna e Robinson somam "
    f"{int(fab.loc['CESSNA', 'n'] + fab.loc['ROBINSON', 'n'])} ocorrências e "
    f"{int(fab.loc['CESSNA', 'a'] + fab.loc['ROBINSON', 'a'])} acidentes. A idade mediana da aeronave é de "
    f"{int(acid.idade_aeronave.median())} anos nos acidentes e de {int(demais.idade_aeronave.median())} anos nas "
    f"demais ocorrências."
)
legenda("Tabela 3.", "Acidentes por tipo de operação, características da aeronave e fase do voo", antes=6, depois=2)
nome_fase = {"SOLO": "Solo", "DECOLAGEM E SUBIDA": "Decolagem e subida", "CRUZEIRO E DESCIDA": "Cruzeiro e descida",
             "APROXIMAÇÃO E POUSO": "Aproximação e pouso"}
linhas = [["Tipo de operação", "", "", ""]]
linhas += [linha_taxa(f"   {rotulo[k]}", r) for k, r in op.iterrows()]
linhas += [["Tipo de motor", "", "", ""]]
linhas += [linha_taxa(f"   {k.capitalize()}", r) for k, r in mot.iterrows()]
linhas += [["Equipamento", "", "", ""]]
linhas += [linha_taxa(f"   {k.capitalize()}", r) for k, r in eqp.iterrows()]
linhas += [["Fase do voo", "", "", ""]]
linhas += [linha_taxa(f"   {nome_fase[k]}", r) for k, r in fase.iterrows()]
tabela(["Atributo", "Ocorrências", "Acidentes", "% acidentes"], linhas, [2.6, 1.15, 1.0, 1.15],
       negrito_primeira_col=("Tipo de operação", "Tipo de motor", "Equipamento", "Fase do voo"))
paragrafo(
    f"Esses atributos, porém, não são independentes. O táxi aéreo reúne {int((taxi.tipo_motor == 'PISTÃO').sum())} "
    f"das {int((b.tipo_motor == 'PISTÃO').sum())} aeronaves a pistão e {int((taxi.tipo_equipamento == 'HELICÓPTERO').sum())} "
    f"dos {int((b.tipo_equipamento == 'HELICÓPTERO').sum())} helicópteros da base, além de aeronaves mais leves e "
    f"mais antigas. Na Regressão Logística, essa correlação torna instáveis as razões de chance de cada atributo "
    f"isolado, que devem ser interpretadas em conjunto ou após a seleção de um subconjunto não redundante; o "
    f"Random Forest é menos sensível a esse problema. Quanto à fase do voo, nenhuma das "
    f"{mil(fase.loc['SOLO', 'n'])} ocorrências em solo resultou em acidente, e as taxas mais altas estão no "
    f"cruzeiro e descida ({pct(fase.loc['CRUZEIRO E DESCIDA', 'a'], fase.loc['CRUZEIRO E DESCIDA', 'n'], 2)}%) e "
    f"na aproximação e pouso ({pct(fase.loc['APROXIMAÇÃO E POUSO', 'a'], fase.loc['APROXIMAÇÃO E POUSO', 'n'], 2)}%)."
)

# ---------------------------------------------------------------- 4.4
subtitulo("4.4. Tipos de ocorrência")
paragrafo(
    f"Os {tipos_unicos} tipos de ocorrência registrados foram agrupados em sete categorias "
    f"operacionais. Nas {int((b.n_tipos > 1).sum())} ocorrências com mais de um tipo, prevaleceu o mais próximo da "
    f"perda da aeronave, nesta ordem de precedência: perda de controle ou colisão em voo, pista, pouso e solo, "
    f"tráfego aéreo, meteorologia, fauna e falha técnica. A Tabela 4 apresenta o resultado."
)
legenda("Tabela 4.", "Ocorrências, incidentes graves e acidentes por categoria de tipo", antes=6, depois=2)
cap = lambda s: s[0] + s[1:].lower()
linhas = [[cap(c), mil(r.n), pct(r.n), int(r.ig), int(r.a), pct(r.a, r.n, 2)] for c, r in cat.iterrows()]
linhas.append(["Total", mil(n), "100,0", len(ig), len(acid), pct(len(acid), n, 2)])
tabela(["Categoria", "Ocorrências", "%", "Incidentes graves", "Acidentes", "% acidentes"], linhas,
       [2.25, 0.95, 0.55, 0.95, 0.8, 0.9], negrito_primeira_col=("Total",))
paragrafo(
    f"Falhas técnicas e colisões com fauna somam {pct(cat.loc[['FALHA TÉCNICA', 'COLISÃO COM FAUNA'], 'n'].sum())}% "
    f"das ocorrências, mas as colisões com fauna quase nunca evoluem para acidente "
    f"({int(cat.loc['COLISÃO COM FAUNA', 'a'])} em {mil(cat.loc['COLISÃO COM FAUNA', 'n'])}). Dentro das falhas "
    f"técnicas, os acidentes concentram-se na falha do motor em voo ({int(motor_voo.Y.sum())} "
    f"de {len(motor_voo)} registros). Perda de controle ou colisão em voo "
    f"({pct(cat.loc['PERDA DE CONTROLE OU COLISÃO EM VOO', 'a'], cat.loc['PERDA DE CONTROLE OU COLISÃO EM VOO', 'n'])}%) "
    f"e pista, pouso e solo ({pct(cat.loc['PISTA, POUSO E SOLO', 'a'], cat.loc['PISTA, POUSO E SOLO', 'n'])}%) têm "
    f"as maiores taxas, e as ocorrências com excursão de pista somam {int(excursao.Y.sum())} acidentes em "
    f"{len(excursao)} registros. "
    f"O tráfego aéreo não registra acidentes, mas reúne {int(cat.loc['TRÁFEGO AÉREO', 'ig'])} incidentes graves, o "
    f"que mostra que incidentes graves e acidentes têm perfis diferentes e que a classe Y = 0 agrupa um "
    f"subconjunto com características próprias."
)

# ---------------------------------------------------------------- 4.5
subtitulo("4.5. Distribuição temporal e geográfica")
fauna23 = int(((b.ano.isin([2023, 2024])) & (b.categoria_tipo == "COLISÃO COM FAUNA")).sum())
paragrafo(
    f"A Figura 1 mostra uma descontinuidade a partir de 2023. De {ini.year} a 2022 há {mil(len(pre))} ocorrências, "
    f"e de 2023 a {fim.year}, {mil(len(pos))}, com pico de {mil((b.ano == 2024).sum())} em 2024. O aumento está quase "
    f"todo na operação regular e nas colisões com fauna ({mil(fauna23)} registros só em 2023 e 2024), enquanto o "
    f"número de acidentes não cresce. Assim, a proporção de acidentes cai de {pct(pre.Y.sum(), len(pre))}% antes "
    f"de 2023 para {pct(pos.Y.sum(), len(pos))}% depois, o que é compatível com uma ampliação da notificação de "
    f"incidentes, e não com uma queda real do risco. Há também defasagem de registro nos meses mais recentes: "
    f"o táxi aéreo tem {int(((b.ano == fim.year - 1) & (b.tipo_operacao == 'TÁXI AÉREO')).sum())} ocorrências em "
    f"{fim.year - 1} e só {int(((b.ano == fim.year) & (b.tipo_operacao == 'TÁXI AÉREO')).sum())} em {fim.year}. "
    f"Por isso, o ano deve ser usado apenas como controle, e não como preditor, e convém avaliar a exclusão dos "
    f"meses finais da janela."
)
figura(FIG_DIR / "fig1_anual.png")
legenda("Figura 1.", f"Ocorrências (acima) e acidentes (abaixo) por ano e tipo de operação; "
                     f"* {ini.year} e {fim.year} parciais")
paragrafo(
    f"Na distribuição geográfica (Figura 2a), o Sudeste concentra {pct(reg.loc['SUDESTE', 'n'])}% das ocorrências "
    f"e tem taxa de acidentes de {pct(reg.loc['SUDESTE', 'a'], reg.loc['SUDESTE', 'n'], 2)}%, enquanto o Norte, "
    f"com {pct(reg.loc['NORTE', 'n'])}% das ocorrências, responde por {int(reg.loc['NORTE', 'a'])} dos "
    f"{len(acid)} acidentes ({pct(reg.loc['NORTE', 'a'], reg.loc['NORTE', 'n'])}%). O contraste também reflete a "
    f"operação: {pct(((b.regiao == 'NORTE') & (b.tipo_operacao == 'TÁXI AÉREO')).sum(), reg.loc['NORTE', 'n'], 0)}% "
    f"das ocorrências do Norte são de táxi aéreo, contra {pct(len(taxi), n, 0)}% no país. O local da ocorrência "
    f"(Figura 2b) repete o padrão: os aeródromos de código SB, em geral os de maior movimento, reúnem "
    f"{pct(loc.loc['AERÓDROMOS SB', 'n'])}% das ocorrências e taxa de "
    f"{pct(loc.loc['AERÓDROMOS SB', 'a'], loc.loc['AERÓDROMOS SB', 'n'], 2)}%, contra "
    f"{pct(loc.loc['FORA DE AERÓDROMO', 'a'], loc.loc['FORA DE AERÓDROMO', 'n'])}% fora de aeródromo e "
    f"{pct(loc.loc['DEMAIS AERÓDROMOS', 'a'], loc.loc['DEMAIS AERÓDROMOS', 'n'])}% nos demais aeródromos. Como há "
    f"{b.aerodromo.nunique()} códigos de aeródromo e {b.cidade.nunique()} municípios, propõe-se usar a região e o "
    f"grupo de local, ou a codificação pelo alvo da Seção 2.2."
)
figura(FIG_DIR / "fig2_regiao_local.png")
legenda("Figura 2.", "Acidentes por 100 ocorrências segundo a região (a) e o local da ocorrência (b); os rótulos "
                     "indicam acidentes/ocorrências")

# ---------------------------------------------------------------- 4.6
subtitulo("4.6. Síntese e implicações para a modelagem")
paragrafo(
    "A base construída contém todos os preditores definidos na Seção 2.2, que são os disponíveis antes da investigação: "
    "tipo de operação, fase do voo, região, fabricante, tipo e quantidade de motores, tipo de equipamento, "
    "idade e peso da aeronave, categoria do tipo de ocorrência e componentes de calendário. Os fatores "
    "contribuintes e as recomendações foram cruzados e estão na base, mas, por serem produzidos pela "
    "investigação, devem ser usados apenas na análise descritiva, e não como preditores. Os principais "
    "cuidados para a etapa preditiva são a forte associação entre o táxi aéreo e as características da "
    "aeronave, a descontinuidade de registro a partir de 2023, a defasagem dos meses mais recentes e o "
    f"desbalanceamento de {pct(len(acid), casas=2)}% na classe positiva, tratado pelas técnicas da Seção 2.2."
)

p = ancora.insert_paragraph_before()
p.add_run().add_break(WD_BREAK.PAGE)
doc.save(DOC_OUT)
print("ok", DOC_OUT)

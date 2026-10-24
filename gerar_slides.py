# -*- coding: utf-8 -*-
"""Slides da apresentacao em video (8-10 min) do relatorio parcial, gerados a partir do artigo.

Uso: python3 gerar_slides.py            (python-pptx; sem dependencias de dados)
Le as tabelas de artigo/tabelas/*.tex e as figuras de artigo/figuras/, para que os numeros dos slides
sejam exatamente os do artigo. Saida: apresentacao/slides_relatorio_parcial.pptx.
Cada slide traz, nas notas do apresentador, o roteiro sugerido e o tempo-alvo.
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

RAIZ = Path(__file__).resolve().parent
FIG = RAIZ / "artigo" / "figuras"
TAB = RAIZ / "artigo" / "tabelas"
SAIDA = RAIZ / "apresentacao" / "slides_relatorio_parcial.pptx"
SAIDA.parent.mkdir(exist_ok=True)

AZUL = RGBColor(0x1F, 0x3B, 0x73)
LARANJA = RGBColor(0xE3, 0x9A, 0x4C)
VERMELHO = RGBColor(0xB2, 0x3A, 0x3A)
CINZA = RGBColor(0x59, 0x59, 0x59)
CINZA_CLARO = RGBColor(0xF2, 0xF2, 0xF2)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
PRETO = RGBColor(0x1A, 0x1A, 0x1A)
FONTE = "Calibri"

LARG, ALT = Inches(13.333), Inches(7.5)
MARGEM = Inches(0.6)
TOPO_CONTEUDO = Inches(1.35)
ALT_CONTEUDO = Inches(5.45)
LARG_CONTEUDO = LARG - 2 * MARGEM

prs = Presentation()
prs.slide_width, prs.slide_height = LARG, ALT
BRANCO_LAYOUT = prs.slide_layouts[6]
TOTAL_SLIDES = 17


# ------------------------------------------------------------------ utilitarios
def tex_limpo(s):
    """Converte um fragmento LaTeX de celula em texto simples; devolve (texto, negrito, italico, recuo)."""
    s = s.strip()
    negrito = "\\textbf{" in s
    italico = "\\textit{" in s
    recuo = "\\hspace" in s
    s = re.sub(r"\\hspace\{[^}]*\}", "", s)
    s = re.sub(r"\\multicolumn\{\d+\}\{[^}]*\}\{(.*)\}$", r"\1", s)
    for _ in range(3):
        s = re.sub(r"\\(textbf|textit)\{([^{}]*)\}", r"\2", s)
    s = s.replace("$\\pm$", "±").replace("$\\dagger$", "†").replace("\\%", "%").replace("--", "–")
    s = re.sub(r"\$n=([\d.]+)\$", r"n = \1", s)
    s = s.replace("$", "").replace("\\\\", "").strip()
    return s, negrito, italico, recuo


def ler_tabela_tex(nome):
    """Le um fragmento tabular e devolve lista de linhas; cada linha e lista de celulas (texto, negrito, italico, recuo)."""
    linhas = []
    for raw in (TAB / nome).read_text(encoding="utf-8").splitlines():
        raw = raw.strip()
        if not raw or raw.startswith("\\begin") or raw.startswith("\\end") or raw.endswith("rule"):
            continue
        raw = raw.rstrip("\\")
        linhas.append([tex_limpo(c) for c in re.split(r"\s*&\s*", raw)])  # linha iniciada por "&" gera 1a celula vazia
    return linhas


def caixa_texto(slide, x, y, w, h, texto, tam=18, cor=PRETO, negrito=False, alinh=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP, italico=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = alinh
    r = p.add_run()
    r.text = texto
    r.font.name, r.font.size, r.font.bold, r.font.italic, r.font.color.rgb = FONTE, Pt(tam), negrito, italico, cor
    return tb


def topicos(slide, x, y, w, h, itens, tam=18, espaco=6):
    """Lista de topicos; cada item pode ser str ou (str, nivel). Trechos entre ** ficam em negrito."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    primeiro = True
    for item in itens:
        texto, nivel = (item, 0) if isinstance(item, str) else item
        p = tf.paragraphs[0] if primeiro else tf.add_paragraph()
        primeiro = False
        p.space_after = Pt(espaco)
        p.level = nivel
        marcador = "•  " if nivel == 0 else "–  "
        partes = texto.split("**")
        for k, parte in enumerate(partes):
            if k == 0:
                parte = (marcador if nivel == 0 else "      " + marcador) + parte
            if not parte:
                continue
            r = p.add_run()
            r.text = parte
            r.font.name, r.font.size, r.font.color.rgb = FONTE, Pt(tam - 2 * nivel), PRETO
            r.font.bold = k % 2 == 1
    return tb


def novo_slide(titulo, parte=None, numero=None):
    slide = prs.slides.add_slide(BRANCO_LAYOUT)
    barra = slide.shapes.add_shape(1, 0, 0, LARG, Inches(0.12))  # 1 = retangulo
    barra.fill.solid(); barra.fill.fore_color.rgb = AZUL; barra.line.fill.background()
    caixa_texto(slide, MARGEM, Inches(0.3), LARG_CONTEUDO, Inches(0.9), titulo, tam=30, cor=AZUL, negrito=True, ancora=MSO_ANCHOR.MIDDLE)
    if parte:
        caixa_texto(slide, LARG - MARGEM - Inches(4.2), Inches(0.18), Inches(4.2), Inches(0.35), parte, tam=11, cor=CINZA, alinh=PP_ALIGN.RIGHT)
    rodape = "Previsão da severidade de ocorrências aeronáuticas na aviação comercial brasileira · Relatório parcial · RP2 EACH-USP"
    caixa_texto(slide, MARGEM, ALT - Inches(0.45), LARG_CONTEUDO - Inches(1), Inches(0.3), rodape, tam=10, cor=CINZA)
    if numero:
        caixa_texto(slide, LARG - MARGEM - Inches(1), ALT - Inches(0.45), Inches(1), Inches(0.3), f"{numero}/{TOTAL_SLIDES}", tam=10, cor=CINZA, alinh=PP_ALIGN.RIGHT)
    return slide


def notas(slide, apresentador, tempo, roteiro):
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"[{apresentador} | tempo-alvo: {tempo}]\n\n{roteiro}"


def figura(slide, nome, x, y, w_max, h_max, px):
    """Insere a figura preservando a proporcao dentro da caixa (x, y, w_max, h_max)."""
    pw, ph = px
    escala = min(w_max / pw, h_max / ph)
    w, h = int(pw * escala), int(ph * escala)
    return slide.shapes.add_picture(str(FIG / nome), x + (w_max - w) // 2, y + (h_max - h) // 2, w, h)


def tabela(slide, linhas, x, y, w, tam=12, larguras=None, alt_linha=Inches(0.3), destaque=()):
    """Tabela nativa a partir das linhas lidas do .tex. Primeira linha = cabecalho."""
    ncol = max(len(l) for l in linhas)
    tb = slide.shapes.add_table(len(linhas), ncol, x, y, w, alt_linha * len(linhas)).table
    if larguras:
        total = sum(larguras)
        for j, frac in enumerate(larguras):
            tb.columns[j].width = int(w * frac / total)
    for i, linha in enumerate(linhas):
        tb.rows[i].height = alt_linha
        grupo = len(linha) == 1 and i > 0
        for j in range(ncol):
            cel = tb.cell(i, j)
            texto, negrito, italico, recuo = linha[j] if j < len(linha) else ("", False, False, False)
            cel.margin_left = cel.margin_right = Inches(0.06)
            cel.margin_top = cel.margin_bottom = Inches(0.02)
            cel.vertical_anchor = MSO_ANCHOR.MIDDLE
            cel.fill.solid()
            if i == 0:
                cel.fill.fore_color.rgb = AZUL
            elif i in destaque:
                cel.fill.fore_color.rgb = RGBColor(0xFD, 0xEB, 0xD3)
            else:
                cel.fill.fore_color.rgb = BRANCO if i % 2 else CINZA_CLARO
            p = cel.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 or grupo else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = ("    " if recuo else "") + texto
            r.font.name, r.font.size = FONTE, Pt(tam)
            r.font.bold = (i == 0) or negrito
            r.font.italic = italico and i > 0
            r.font.color.rgb = BRANCO if i == 0 else PRETO
        if grupo:
            tb.cell(i, 0).merge(tb.cell(i, ncol - 1))
    return tb


def legenda(slide, texto, y=None):
    caixa_texto(slide, MARGEM, y or (ALT - Inches(0.95)), LARG_CONTEUDO, Inches(0.4), texto, tam=12, cor=CINZA, italico=True, alinh=PP_ALIGN.CENTER)


# ================================================================== SLIDE 1: titulo
s = prs.slides.add_slide(BRANCO_LAYOUT)
faixa = s.shapes.add_shape(1, 0, 0, LARG, Inches(4.2)); faixa.fill.solid(); faixa.fill.fore_color.rgb = AZUL; faixa.line.fill.background()
caixa_texto(s, MARGEM, Inches(1.1), LARG_CONTEUDO, Inches(2.0), "Previsão da severidade de ocorrências aeronáuticas na aviação comercial brasileira (RBAC 121 e 135)", tam=34, cor=BRANCO, negrito=True, ancora=MSO_ANCHOR.MIDDLE)
caixa_texto(s, MARGEM, Inches(3.1), LARG_CONTEUDO, Inches(0.6), "Relatório parcial · RP2 · 2026", tam=18, cor=RGBColor(0xC9, 0xD6, 0xEE))
caixa_texto(s, MARGEM, Inches(4.6), LARG_CONTEUDO, Inches(0.6), "Eric I. W. Chou · João Fontanezi · Marcos R. da S. Santos", tam=22, cor=PRETO, negrito=True)
caixa_texto(s, MARGEM, Inches(5.2), LARG_CONTEUDO, Inches(0.5), "Escola de Artes, Ciências e Humanidades – Universidade de São Paulo (EACH-USP)", tam=16, cor=CINZA)
caixa_texto(s, MARGEM, Inches(6.2), LARG_CONTEUDO, Inches(0.5), "Dados abertos do CENIPA / Painel SIPAER · Regressão Logística e Random Forest · github.com/Fontanezi/RP2_ARTIGO_2026", tam=13, cor=CINZA)
notas(s, "Todos (abertura)", "0:15", "Apresentar o título e os integrantes. Uma frase: 'a pergunta é se dá para antecipar, com o que se sabe no momento da notificação, se uma ocorrência é um incidente, um incidente grave ou um acidente'.")

# ================================================================== PARTE 1
s = novo_slide("Contexto e problema", "Parte 1 · Introdução", 2)
topicos(s, MARGEM, TOPO_CONTEUDO, LARG_CONTEUDO, ALT_CONTEUDO, [
    "Transporte aéreo público no Brasil: **RBAC 121** (aviões com mais de 19 assentos) e **RBAC 135** (até 19 assentos ou helicópteros: táxi aéreo e operação não regular)",
    "O **CENIPA** investiga e publica todas as ocorrências no Painel SIPAER e em dados.gov.br",
    "A análise usual desses registros é **retrospectiva e caso a caso**: explica o que aconteceu, não antecipa",
    "A literatura nacional traz séries temporais agregadas e estudos descritivos, mas **não classifica a severidade de uma ocorrência individual**",
    "Pergunta do trabalho: com os atributos conhecidos no momento da notificação, é possível estimar a probabilidade de a ocorrência ser **incidente, incidente grave ou acidente**?",
], tam=19, espaco=10)
notas(s, "Integrante 1", "0:35", "Situar o setor (RBAC 121 e 135), a fonte (CENIPA) e a lacuna: análise retrospectiva, nenhum modelo que classifique a severidade de um evento individual. Terminar enunciando a pergunta do trabalho.")

s = novo_slide("Objetivo e as três classes do CENIPA (NSCA 3-13)", "Parte 1 · Introdução", 3)
topicos(s, MARGEM, TOPO_CONTEUDO, LARG_CONTEUDO, Inches(1.1), [
    "**Objetivo:** desenvolver e comparar modelos de Aprendizado de Máquina (Regressão Logística e Random Forest) para classificar a severidade de ocorrências da aviação comercial, usando apenas atributos disponíveis antes da investigação",
], tam=18)
y0 = TOPO_CONTEUDO + Inches(1.3)
for k, (titulo, cor, texto) in enumerate([
    ("Acidente", VERMELHO, "Lesão grave ou morte, dano ou falha estrutural que afete a resistência ou o desempenho da aeronave (ou exija grande reparo), aeronave desaparecida ou inacessível. Exceções: danos localizados (um motor, hélices, pneus, colisão com fauna)."),
    ("Incidente grave", LARANJA, "Incidente cujas circunstâncias indicam elevado potencial de risco de acidente. A diferença para o acidente está apenas nas consequências: é o quase-acidente."),
    ("Incidente", AZUL, "Toda ocorrência que não se caracteriza como acidente, mas que afeta ou pode afetar a segurança da operação."),
]):
    x = MARGEM + k * (LARG_CONTEUDO // 3)
    w = LARG_CONTEUDO // 3 - Inches(0.2)
    cab = s.shapes.add_shape(1, x, y0, w, Inches(0.55)); cab.fill.solid(); cab.fill.fore_color.rgb = cor; cab.line.fill.background()
    cab.text_frame.text = titulo
    pr = cab.text_frame.paragraphs[0]; pr.alignment = PP_ALIGN.CENTER
    pr.runs[0].font.name, pr.runs[0].font.size, pr.runs[0].font.bold, pr.runs[0].font.color.rgb = FONTE, Pt(18), True, BRANCO
    corpo = s.shapes.add_shape(1, x, y0 + Inches(0.55), w, Inches(2.6)); corpo.fill.solid(); corpo.fill.fore_color.rgb = CINZA_CLARO; corpo.line.fill.background()
    caixa_texto(s, x + Inches(0.1), y0 + Inches(0.65), w - Inches(0.2), Inches(2.4), texto, tam=14)
caixa_texto(s, MARGEM, y0 + Inches(3.35), LARG_CONTEUDO, Inches(0.5), "Por que três classes e não duas? As definições separam o incidente grave pelo potencial de risco, não pela consequência; e os dados mostram perfis diferentes (slide 9).", tam=14, cor=CINZA, italico=True)
notas(s, "Integrante 1", "0:40", "Ler o objetivo. Depois as três definições, com ênfase em 'a diferença entre incidente grave e acidente está apenas nas consequências'. Fechar com a justificativa de manter as três classes separadas em vez de agrupar incidente grave com incidente.")

s = novo_slide("Trabalhos correlatos", "Parte 1 · Introdução", 4)
corr = [[("Referência", 1, 0, 0), ("Modelo", 1, 0, 0), ("Dados utilizados", 1, 0, 0), ("Contribuição deste trabalho", 1, 0, 0)],
        [("Reinert (2023)", 0, 0, 0), ("Ajuste de curvas e ARIMA sobre séries anuais agregadas", 0, 0, 0), ("SIPAER 2012–2022, aviões Ipanema e Sêneca", 0, 0, 0), ("Classifica cada ocorrência individual em três classes com atributos pré-investigação; cobre todo o transporte aéreo comercial", 0, 0, 0)],
        [("Soares (2012)", 0, 0, 0), ("Análise sinótica qualitativa de casos", 0, 0, 0), ("Acidentes dos anos 1990 associados à meteorologia", 0, 0, 0), ("Quantifica o risco com modelos discriminativos validados por métricas; meteorologia é uma das categorias de tipo", 0, 0, 0)],
        [("Barros e Balero (2012)", 0, 0, 0), ("Pesquisa bibliográfica e estatística descritiva", 0, 0, 0), ("Estações meteorológicas do Centro-Oeste e estatísticas do SIPAER", 0, 0, 0), ("Abrangência nacional com a região como preditor; avaliação fora da amostra com baselines explícitos", 0, 0, 0)]]
tabela(s, corr, MARGEM, TOPO_CONTEUDO, LARG_CONTEUDO, tam=14, larguras=[1.6, 2.6, 2.8, 4.2], alt_linha=Inches(0.95))
caixa_texto(s, MARGEM, TOPO_CONTEUDO + Inches(4.1), LARG_CONTEUDO, Inches(0.8), "Em comum: nenhum dos trabalhos estima a severidade de uma ocorrência individual a partir de atributos operacionais, geográficos e técnicos. É essa a lacuna que este trabalho ataca.", tam=15, cor=AZUL, negrito=True)
notas(s, "Integrante 1", "0:35", "Um trabalho de séries temporais (Reinert) e dois descritivos sobre meteorologia (Soares; Barros e Balero). Destacar a última coluna: o que este trabalho faz que eles não fazem. Não ler a tabela inteira.")

# ================================================================== PARTE 2
s = novo_slide("Dados: construção da base e escopo", "Parte 2 · Metodologia", 5)
etapas = [("14.626", "ocorrências nos arquivos do CENIPA"), ("9.471", "entre 10/05/2016 e 09/05/2026"), ("5.038", "com aeronave em operação regular, não regular ou táxi aéreo")]
for k, (num, desc) in enumerate(etapas):
    x = MARGEM + k * Inches(4.1)
    bloco = s.shapes.add_shape(1, x, TOPO_CONTEUDO, Inches(3.7), Inches(1.5)); bloco.fill.solid(); bloco.fill.fore_color.rgb = AZUL if k == 2 else CINZA_CLARO; bloco.line.fill.background()
    caixa_texto(s, x, TOPO_CONTEUDO + Inches(0.1), Inches(3.7), Inches(0.7), num, tam=32, cor=BRANCO if k == 2 else AZUL, negrito=True, alinh=PP_ALIGN.CENTER)
    caixa_texto(s, x + Inches(0.1), TOPO_CONTEUDO + Inches(0.8), Inches(3.5), Inches(0.65), desc, tam=13, cor=BRANCO if k == 2 else CINZA, alinh=PP_ALIGN.CENTER)
    if k < 2:
        caixa_texto(s, x + Inches(3.7), TOPO_CONTEUDO + Inches(0.4), Inches(0.4), Inches(0.7), "→", tam=28, cor=CINZA, alinh=PP_ALIGN.CENTER)
topicos(s, MARGEM, TOPO_CONTEUDO + Inches(1.8), LARG_CONTEUDO, Inches(3.6), [
    "Cruzamento das tabelas de Ocorrências, Tipos, Aeronaves, Fatores Contribuintes e Recomendações pelo código da ocorrência; **uma linha por ocorrência**",
    "Classes: **4.845 incidentes (96,2%) · 102 incidentes graves (2,0%) · 91 acidentes (1,8%)**; 17 acidentes com vítimas fatais, 122 mortes (62 em Vinhedo, 2024)",
    "Escopo = todo o transporte aéreo comercial (RBAC 121 + 135). Por quê? Mesma família regulatória; a operação regular sozinha tem só **12 acidentes**, insuficiente para aprender; e o tipo de operação entra como preditor",
    "Base publicada com dicionário de dados: 5.038 linhas × 44 colunas (BASE_SIPAER.xlsx)",
], tam=17, espaco=8)
notas(s, "Integrante 2", "0:35", "Funil de construção da base (três números). Depois o desbalanceamento: menos de 2% de acidentes. Justificar o escopo com os três motivos, em especial os 12 acidentes da operação regular.")

s = novo_slide("Vazamento de informação: o que não pode ser preditor", "Parte 2 · Metodologia", 6)
pos = ler_tabela_tex("tab_posevento.tex")
tabela(s, pos, MARGEM, TOPO_CONTEUDO, Inches(7.4), tam=12, larguras=[3.9, 1.4, 1.5, 1.6], alt_linha=Inches(0.45))
caixa_texto(s, MARGEM, TOPO_CONTEUDO + Inches(2.85), Inches(7.4), Inches(0.6), "Atributos produzidos pela investigação, por classe. A investigação completa é feita quase só em acidentes e incidentes graves: esses atributos entregariam a resposta ao modelo.", tam=12, cor=CINZA, italico=True)
topicos(s, MARGEM + Inches(7.7), TOPO_CONTEUDO, LARG_CONTEUDO - Inches(7.7), ALT_CONTEUDO, [
    "**Fora dos preditores:** nível de dano, fatalidades, situação da investigação, relatório final, fatores contribuintes, recomendações; e o ano (mudança de cobertura em 2023)",
    "**Preditores (disponíveis na notificação):** tipo de operação, equipamento, tipo e quantidade de motores, fabricante, peso máximo de decolagem, assentos, idade, fase do voo, região, local, categoria do tipo de ocorrência, saída de pista, aeronaves envolvidas, mês, dia da semana, hora",
    "Aeródromo (203 níveis) e tipo detalhado (65) entram agrupados",
], tam=15, espaco=10)
notas(s, "Integrante 2", "0:40", "Este é o achado metodológico central. Mostrar na tabela que relatório, fatores e recomendações existem em 61–68% dos acidentes, 44–45% dos incidentes graves e 0,1% dos incidentes: usar isso seria trapaça. Listar rapidamente o que ficou como preditor.")

s = novo_slide("Pipeline e protocolo de avaliação", "Parte 2 · Metodologia", 7)
caixas = ["Imputação\n(mediana / moda)", "SMOTENC\n(só no treino)", "Padronização\n+ one-hot", "Regressão Logística\nmultinomial  ·  Random Forest"]
for k, txt in enumerate(caixas):
    x = MARGEM + k * Inches(3.1)
    w = Inches(2.7) if k < 3 else Inches(3.1)
    cx = s.shapes.add_shape(1, x, TOPO_CONTEUDO, w, Inches(1.0)); cx.fill.solid(); cx.fill.fore_color.rgb = AZUL if k == 3 else CINZA_CLARO; cx.line.fill.background()
    caixa_texto(s, x, TOPO_CONTEUDO, w, Inches(1.0), txt, tam=14, cor=BRANCO if k == 3 else PRETO, negrito=k == 3, alinh=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
    if k < 3:
        caixa_texto(s, x + w, TOPO_CONTEUDO + Inches(0.15), Inches(0.4), Inches(0.7), "→", tam=26, cor=CINZA, alinh=PP_ALIGN.CENTER)
caixa_texto(s, MARGEM, TOPO_CONTEUDO + Inches(1.05), LARG_CONTEUDO, Inches(0.4), "Tudo ajustado apenas nas dobras de treino, nunca nos dados avaliados.", tam=12, cor=CINZA, italico=True, alinh=PP_ALIGN.CENTER)
topicos(s, MARGEM, TOPO_CONTEUDO + Inches(1.55), Inches(6.2), Inches(3.8), [
    "**Partição 80/20 estratificada:** treino 4.030 (3.876 / 81 / 73), teste 1.008 (969 / 21 / 18)",
    "**Hiperparâmetros:** busca em grade com validação cruzada estratificada (k = 5) no treino, maximizando F1-macro",
    "**Resultado principal:** predições out-of-fold no treino (73 acidentes); teste (18 acidentes) como confirmação",
    "**Balanceamento, 4 estratégias:** nenhum · pesos de classe · SMOTENC · SMOTENC + pesos",
], tam=15, espaco=8)
topicos(s, MARGEM + Inches(6.5), TOPO_CONTEUDO + Inches(1.55), LARG_CONTEUDO - Inches(6.5), Inches(3.8), [
    "**Métricas:** sensibilidade, precisão e F1 por classe; F1-macro; acurácia balanceada; AUC-ROC e precisão média (um-contra-resto); matriz de confusão 3×3",
    "Acurácia só como referência: prever sempre 'incidente' acerta 96,2%",
    "**Dois baselines:** classe majoritária; regra 'táxi aéreo ou não regular ⇒ acidente'",
    "**Interpretação:** razões de chance (LR) e importância por permutação (RF)",
], tam=15, espaco=8)
notas(s, "Integrante 2", "0:45", "Percorrer o fluxo da esquerda para a direita. Explicar por que o número principal é a validação cruzada (18 acidentes no teste: cada erro move a sensibilidade em 5,6 pontos). Dizer que F1-macro é o critério de seleção e que acurácia não serve. Apresentar os dois baselines: todo modelo tem que bater a regra do tipo de operação.")

# ================================================================== PARTE 3: RESULTADOS (>= 4 min)
s = novo_slide("Resultados · Operação e aeronave separam as classes", "Parte 3 · Resultados parciais", 8)
op = ler_tabela_tex("tab_operacao.tex")
tabela(s, op, MARGEM, TOPO_CONTEUDO, Inches(7.6), tam=12, larguras=[2.6, 1.3, 1.2, 1.2, 1.4, 1.4], alt_linha=Inches(0.285))
topicos(s, MARGEM + Inches(7.9), TOPO_CONTEUDO, LARG_CONTEUDO - Inches(7.9), ALT_CONTEUDO, [
    "Operação regular: 85% das ocorrências, **0,28%** de acidentes",
    "Táxi aéreo: 13% das ocorrências, **72 dos 91 acidentes** (11,1%)",
    "Pistão 22,6% vs. jato 0,32%; helicóptero 7,5% vs. avião 1,3%",
    "Idade mediana: 11 anos (incidente), 16 (inc. grave), **32 (acidente)**; peso mediano: 70,5 t, 7,3 t, **2,2 t**",
    "Nenhum acidente em solo; incidentes graves concentram-se na aproximação e pouso",
    "Atenção: táxi aéreo, pistão, helicóptero e aeronave leve e antiga são **o mesmo subgrupo** (colinearidade)",
], tam=14, espaco=7)
notas(s, "Integrante 3", "0:30", "O tipo de operação é o atributo que mais separa as classes: regular 0,28% de acidentes, táxi aéreo 11%. Aeronaves a pistão, monomotoras e helicópteros seguem o padrão. Avisar da colinearidade: esses atributos descrevem o mesmo subgrupo, o que vai aparecer na interpretação da Regressão Logística.")

s = novo_slide("Resultados · Tipos de ocorrência: perfis distintos", "Parte 3 · Resultados parciais", 9)
tip = ler_tabela_tex("tab_tipos.tex")
tabela(s, tip, MARGEM, TOPO_CONTEUDO, Inches(8.4), tam=13, larguras=[3.2, 1.3, 0.9, 1.3, 1.3, 1.2, 1.1], alt_linha=Inches(0.4))
topicos(s, MARGEM + Inches(8.7), TOPO_CONTEUDO, LARG_CONTEUDO - Inches(8.7), ALT_CONTEUDO, [
    "Falha técnica + fauna = 85% das ocorrências; **colisão com fauna quase nunca é grave** (0 e 1 em 1.767)",
    "**Acidente:** falha do motor em voo (14 acidentes, 0 inc. graves em 122) e perda de controle (60%)",
    "**Incidente grave:** tráfego aéreo (21 inc. graves, **0 acidentes**) e pista, pouso e solo",
    "Perfis diferentes ⇒ faz sentido manter as três classes separadas",
], tam=14, espaco=9)
notas(s, "Integrante 3", "0:30", "Mostrar que incidente grave e acidente vêm de tipos diferentes: tráfego aéreo gera 21 incidentes graves e nenhum acidente; falha do motor em voo gera 14 acidentes e nenhum incidente grave. Esse é o argumento empírico das três classes.")

s = novo_slide("Resultados · Distribuição temporal: a cobertura mudou em 2023", "Parte 3 · Resultados parciais", 10)
figura(s, "fig1_anual.png", MARGEM, TOPO_CONTEUDO, Inches(7.6), ALT_CONTEUDO - Inches(0.1), (1584, 946))
topicos(s, MARGEM + Inches(7.9), TOPO_CONTEUDO, LARG_CONTEUDO - Inches(7.9), ALT_CONTEUDO, [
    "2016–2022: 1.097 ocorrências; 2023–2026: **3.941** (pico de 1.582 em 2024)",
    "O salto está na operação regular e nas colisões com fauna (1.462 só em 2023–24)",
    "Acidentes e incidentes graves por ano **não crescem**",
    "Proporção de acidentes cai de 5,9% para 0,7%: mais notificação de incidentes, não menos risco",
    "Por isso o **ano não é preditor** (só controle)",
], tam=14, espaco=9)
legenda(s, "Figura 1. Ocorrências por ano e tipo de operação (acima) e incidentes graves e acidentes por ano (abaixo); 2016 e 2026 parciais", y=ALT - Inches(0.85))
notas(s, "Integrante 3", "0:30", "Apontar a descontinuidade em 2023 no painel de cima e a estabilidade dos acidentes no painel de baixo. Conclusão: é mudança de cobertura das notificações, e o ano não pode ser preditor.")

s = novo_slide("Resultados · Distribuição geográfica", "Parte 3 · Resultados parciais", 11)
figura(s, "fig2_regiao_local.png", MARGEM, TOPO_CONTEUDO, Inches(8.4), Inches(3.5), (1584, 638))
topicos(s, MARGEM + Inches(8.7), TOPO_CONTEUDO, LARG_CONTEUDO - Inches(8.7), ALT_CONTEUDO, [
    "**Norte:** 10,5% das ocorrências, **44 dos 91 acidentes** (8,3%); 35% das suas ocorrências são táxi aéreo (13% no país)",
    "Sudeste: metade das ocorrências, 0,8% de acidentes",
    "**Local separa as classes graves:** aeródromos SB têm 57 dos 102 incidentes graves, mas só 0,33% de acidentes; **49 acidentes ocorreram fora de aeródromo**",
    "Preditores usados: região e grupo de local (203 aeródromos e 241 municípios são demais para 91 acidentes)",
], tam=14, espaco=9)
legenda(s, "Figura 2. Incidentes graves e acidentes por 100 ocorrências segundo a região (a) e o local da ocorrência (b)", y=TOPO_CONTEUDO + Inches(3.6))
notas(s, "Integrante 3", "0:25", "Norte concentra acidentes, e isso reflete a operação (táxi aéreo). O local separa as duas classes graves: incidentes graves acontecem nos grandes aeródromos, acidentes fora de aeródromo.")

s = novo_slide("Resultados · Balanceamento: o efeito depende do algoritmo", "Parte 3 · Modelos", 12)
abl = ler_tabela_tex("tab_ablacao.tex")
tabela(s, abl, MARGEM, TOPO_CONTEUDO, LARG_CONTEUDO, tam=13, larguras=[2.3, 2.0, 1.8, 1.1, 1.1, 1.1, 1.1, 1.1], alt_linha=Inches(0.36), destaque=(1, 7))
topicos(s, MARGEM, TOPO_CONTEUDO + Inches(3.4), LARG_CONTEUDO, Inches(1.9), [
    "**Regressão Logística:** sem balanceamento tem o maior F1-macro (0,556), mas detecta só 37% dos acidentes; com pesos detecta 71%, mas a precisão cai para 18%",
    "**Random Forest:** pesos e SMOTENC empatam (0,62) e levam a sensibilidade de 41% para 60–63% mantendo precisão de ~45%; SMOTENC escolhido pela menor variância",
    "**SMOTENC + pesos = SMOTENC:** após igualar as classes, os pesos balanceados ficam todos iguais a 1 (verificação da redundância)",
], tam=13, espaco=6)
legenda(s, "Tabela 5. F1-macro na validação cruzada (média ± desvio-padrão) e métricas out-of-fold, em %; † configuração selecionada; AC = acidente, IG = incidente grave", y=ALT - Inches(0.78))
notas(s, "Integrante 1 ou 2", "0:30", "Ler só as linhas destacadas (†). Explicar a troca sensibilidade × precisão na Regressão Logística com pesos. No Random Forest o balanceamento ajuda sem destruir a precisão. Mencionar que SMOTENC + pesos é idêntico ao SMOTENC, e por quê.")

s = novo_slide("Resultados · Desempenho dos modelos e dos baselines", "Parte 3 · Modelos", 13)
met = ler_tabela_tex("tab_metricas.tex")
tabela(s, met, MARGEM, TOPO_CONTEUDO, LARG_CONTEUDO, tam=12, larguras=[3.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.1, 1.1, 1.0], alt_linha=Inches(0.3), destaque=(5, 10))
topicos(s, MARGEM, TOPO_CONTEUDO + Inches(3.45), LARG_CONTEUDO, Inches(1.85), [
    "**Baselines:** a classe majoritária acerta 96,2% e não detecta nenhuma ocorrência grave; a regra pelo tipo de operação acha 86% dos acidentes com precisão de 10,7% (F1 = 0,19) e nenhum incidente grave",
    "**Random Forest + SMOTENC:** 44 dos 73 acidentes (60%) com precisão de 46% e 35 dos 81 incidentes graves (43%); F1-macro **0,624**. Regressão Logística: 37% e 21%, F1-macro 0,563",
    "**Teste confirma a ordem:** RF 11/18 acidentes e 8/21 inc. graves (F1-macro 0,609); LR 10 e 2 (0,560). AUC acidente no teste 0,97–0,98 (variância de 18 casos)",
], tam=13, espaco=6)
legenda(s, "Tabela 6. Predições out-of-fold no treino (73 acidentes, 81 inc. graves) e teste (18 e 21); sensibilidade e precisão em %; AC = acidente, IG = incidente grave, INC = incidente", y=ALT - Inches(0.78))
notas(s, "Integrante 1 ou 2", "0:45", "Este é o slide central dos resultados. Primeiro os baselines (acurácia enganosa; regra com alta sensibilidade e precisão baixíssima). Depois o Random Forest, linha destacada: 60% dos acidentes com 46% de precisão e 43% dos incidentes graves. Fechar com o teste confirmando a ordem.")

s = novo_slide("Resultados · Onde os modelos erram", "Parte 3 · Modelos", 14)
figura(s, "fig3_confusao.png", MARGEM, TOPO_CONTEUDO, Inches(8.6), Inches(3.6), (1584, 660))
topicos(s, MARGEM + Inches(8.9), TOPO_CONTEUDO, LARG_CONTEUDO - Inches(8.9), ALT_CONTEUDO, [
    "**Random Forest:** acidentes perdidos vão para 'incidente' (23) e raramente para 'inc. grave' (6); 42 incidentes previstos como acidente e 65 como graves (2,8% dos incidentes)",
    "Incidentes graves não detectados dividem-se entre incidente (36) e acidente (10)",
    "**Regressão Logística** quase não prevê 'incidente grave' (35 predições em 4.030): é a classe de sinal mais fraco",
], tam=14, espaco=9)
legenda(s, "Figura 3. Matrizes de confusão das predições out-of-fold no treino (contagem e percentual da linha)", y=TOPO_CONTEUDO + Inches(3.7))
notas(s, "Integrante 1 ou 2", "0:30", "Ler a diagonal do Random Forest (97%, 43%, 60%) e explicar o custo: 107 incidentes classificados acima do que são, 2,8% deles. Na Regressão Logística a coluna 'inc. grave' está quase vazia.")

s = novo_slide("Resultados · Independente do limiar: ROC e precisão-sensibilidade", "Parte 3 · Modelos", 15)
figura(s, "fig4_curvas.png", MARGEM, TOPO_CONTEUDO, Inches(8.6), Inches(3.6), (1584, 660))
topicos(s, MARGEM + Inches(8.9), TOPO_CONTEUDO, LARG_CONTEUDO - Inches(8.9), ALT_CONTEUDO, [
    "**Acidente:** AUC 0,94 nos dois modelos; precisão média 0,41–0,44, cerca de **24× a prevalência** (1,8%)",
    "**Incidente grave:** AUC 0,85–0,88, precisão média 0,23–0,33: a classe mais difícil",
    "O losango é a regra pelo tipo de operação. **Na mesma sensibilidade (86%)**, os modelos chegam a 14–17% de precisão contra 10,7%: ganho real, mas modesto",
    "A vantagem está na faixa intermediária: RF com 45–70% de precisão para 40–65% de sensibilidade",
], tam=13, espaco=8)
legenda(s, "Figura 4. Curvas ROC (a) e precisão-sensibilidade (b) um-contra-resto, predições out-of-fold; o losango marca a regra pelo tipo de operação", y=TOPO_CONTEUDO + Inches(3.7))
notas(s, "Integrante 1 ou 2", "0:30", "As curvas mostram o desempenho sem depender do limiar de decisão. Ponto honesto: na sensibilidade da regra, o ganho de precisão é pequeno; o modelo vale mais na faixa intermediária. O ajuste de limiar é o próximo passo.")

s = novo_slide("Resultados · O que os modelos aprenderam", "Parte 3 · Modelos", 16)
figura(s, "fig5_importancias.png", MARGEM, TOPO_CONTEUDO, Inches(8.2), Inches(4.1), (1584, 792))
topicos(s, MARGEM + Inches(8.5), TOPO_CONTEUDO, LARG_CONTEUDO - Inches(8.5), ALT_CONTEUDO, [
    "**Random Forest:** tipo de operação, peso da aeronave, fase do voo, local e categoria do tipo são os mais importantes; mês, dia, hora e assentos têm importância ~0",
    "**Regressão Logística** (chance de acidente vs. incidente): perda de controle em voo (~10³, só 20 casos), Robinson (58×), turboeixo (48×), fora de aeródromo (26×), táxi aéreo e não regular (15×), Norte (13×)",
    "Cessna, Sikorsky e Embraer aparecem com chance **menor** que Airbus apesar de taxas maiores: é a **colinearidade** com operação, motor e peso; coeficientes devem ser lidos em conjunto",
], tam=13, espaco=9)
legenda(s, "Figura 5. (a) Contrastes de log da razão de chances (acidente vs. incidente) na Regressão Logística; (b) importância por permutação no Random Forest (queda da AUC no teste)", y=TOPO_CONTEUDO + Inches(4.2))
notas(s, "Integrante 1 ou 2", "0:35", "Random Forest confirma a análise exploratória: operação, peso, fase, local, tipo. Calendário não serve. Na Regressão Logística, mostrar os maiores contrastes e explicar o sinal 'errado' de Cessna/Sikorsky/Embraer pela colinearidade: a taxa marginal deles já está explicada por operação, motor e peso.")

# ================================================================== CONCLUSAO
s = novo_slide("Discussão, limitações e próximos passos", "Conclusão", 17)
topicos(s, MARGEM, TOPO_CONTEUDO, Inches(6.1), ALT_CONTEUDO, [
    "**O que foi feito:** base de 5.038 ocorrências; alvo em três classes; exclusão dos atributos pós-investigação; dois algoritmos × quatro balanceamentos",
    "**Melhor modelo:** Random Forest + SMOTENC, F1-macro 0,62; 60% dos acidentes e 43% dos incidentes graves fora da amostra, confirmado no teste",
    "**Regressão Logística:** 0,56; fraca para incidente grave, útil para interpretar",
    "**Leitura honesta:** ambos batem os baselines em F1, mas na sensibilidade da regra o ganho de precisão é modesto; incidente grave é a classe difícil (AUC 0,85–0,88)",
], tam=15, espaco=9)
topicos(s, MARGEM + Inches(6.4), TOPO_CONTEUDO, LARG_CONTEUDO - Inches(6.4), ALT_CONTEUDO, [
    "**Limitações:** 73 + 18 acidentes ⇒ variância alta; categoria do tipo descreve parte do desfecho; cobertura mudou em 2023",
    "**Próximos passos até o relatório final:**",
    ("ajustar o limiar de decisão por custo (sensibilidade ≥ 80% para acidentes)", 1),
    ("remover atributos de calendário; testar codificação pelo alvo de aeródromo e tipo", 1),
    ("ablação da categoria do tipo; validação por período (treino ≤ 2022, teste ≥ 2023)", 1),
    ("intervalos de confiança por validação cruzada repetida; calibração das probabilidades", 1),
    "Código, base e artigo: **github.com/Fontanezi/RP2_ARTIGO_2026**",
], tam=15, espaco=7)
notas(s, "Integrante 3 (fechamento)", "0:40", "Resumir em quatro frases o que foi feito e o melhor resultado. Declarar as limitações antes que perguntem. Listar os próximos passos até 18/11. Encerrar agradecendo.")

assert len(prs.slides) == TOTAL_SLIDES, len(prs.slides)
prs.save(SAIDA)
print("ok", SAIDA, len(prs.slides), "slides")
